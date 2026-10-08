import io
import json
import threading
import time

import pytest
from fastapi.testclient import TestClient

from conftest import PAIR, FakeBackend, FakeTokenizer, thresholds, words
from squint.app import create_app
from squint.config import Limits


def make_client(settings, *, backend=None, factory=None, password=None, tick=0.02, **over):
    if over:
        settings = settings.with_overrides(**over)
    backends = []

    def default_factory():
        b = backend if backend is not None else FakeBackend()
        backends.append(b)
        return b

    app = create_app(settings, backend_factory=factory or default_factory, tokenizer=FakeTokenizer(),
                     password=password, job_tick_s=tick)
    client = TestClient(app)
    client.backends = backends
    client.svc = app.state.svc
    return client


def wait_job(client, job_id, statuses=("done", "error"), timeout=10.0, headers=None):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        r = client.get(f"/api/jobs/{job_id}", headers=headers)
        assert r.status_code == 200, r.text
        snap = r.json()
        if snap["status"] in statuses:
            return snap
        time.sleep(0.01)
    raise AssertionError(f"job {job_id} never reached {statuses}: {snap}")


def wait_for(cond, timeout=5.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if cond():
            return True
        time.sleep(0.01)
    return False


def analyze(client, text, **kw):
    return client.post("/api/analyze", json={"text": text}, **kw)


# ------------------------------------------------------------------------------------- session


def test_no_password_mode(settings):
    with make_client(settings) as c:
        assert c.get("/api/session").json() == {"authenticated": True, "password_required": False}
        assert c.post("/api/login", json={"password": "anything"}).status_code == 204
        assert c.get("/api/status").status_code == 200


def test_password_flow(settings):
    with make_client(settings, password="hunter2") as c:
        assert c.get("/api/session").json() == {"authenticated": False, "password_required": True}
        for path in ("/api/status", "/api/accuracy", "/api/jobs/j_x"):
            r = c.get(path)
            assert r.status_code == 401 and r.json()["error"]["code"] == "auth_required"
        assert analyze(c, words(30)).status_code == 401
        assert c.post("/api/extract", files={"file": ("a.txt", b"hello there")}).status_code == 401

        r = c.post("/api/login", json={"password": "wrong"})
        assert r.status_code == 401 and r.json()["error"]["code"] == "bad_password"
        r = c.post("/api/login", json={"password": "hunter2"})
        assert r.status_code == 204
        cookie = r.headers["set-cookie"].lower()
        assert "squint_session=" in cookie and "httponly" in cookie and "samesite=lax" in cookie
        assert "secure" not in cookie  # plain http on localhost
        assert c.get("/api/session").json()["authenticated"] is True
        assert c.get("/api/status").status_code == 200

        token = c.cookies.get("squint_session")
        c.cookies.set("squint_session", token[:-2] + ("AA" if token[-2:] != "AA" else "BB"))
        assert c.get("/api/status").status_code == 401
        c.cookies.set("squint_session", token)
        assert c.post("/api/logout").status_code == 204
        c.cookies.clear()
        assert c.get("/api/session").json()["authenticated"] is False


def test_cookie_is_secure_behind_https_proxy(settings):
    with make_client(settings, password="pw") as c:
        r = c.post("/api/login", json={"password": "pw"}, headers={"X-Forwarded-Proto": "https"})
        assert "secure" in r.headers["set-cookie"].lower()


def test_login_attempts_are_limited(settings):
    with make_client(settings, password="pw", login_max_failures=3) as c:
        ip = {"CF-Connecting-IP": "203.0.113.5"}
        for _ in range(3):
            assert c.post("/api/login", json={"password": "nope"}, headers=ip).status_code == 401
        r = c.post("/api/login", json={"password": "pw"}, headers=ip)  # even the right one
        assert r.status_code == 429
        body = r.json()["error"]
        assert body["code"] == "rate_limited" and body["retry_after_s"] > 0
        assert int(r.headers["retry-after"]) == body["retry_after_s"]
        # another client is unaffected
        assert c.post("/api/login", json={"password": "pw"}, headers={"CF-Connecting-IP": "198.51.100.7"}).status_code == 204


# ------------------------------------------------------------------------------------- analyze


def test_word_limits(settings):
    with make_client(settings) as c:
        r = analyze(c, words(19))
        assert r.status_code == 422 and r.json()["error"]["code"] == "too_short"
        assert analyze(c, words(20)).status_code == 202
        r = analyze(c, words(2501))
        assert r.status_code == 422 and r.json()["error"]["code"] == "too_long"
        assert analyze(c, words(2500)).status_code == 202
        r = analyze(c, " ".join(["x" * 2000] * 25))  # few words, too many characters
        assert r.status_code == 422 and r.json()["error"]["code"] == "too_long"


def test_bad_requests(settings):
    with make_client(settings) as c:
        r = c.post("/api/analyze", content=b"text=hello", headers={"Content-Type": "application/x-www-form-urlencoded"})
        assert r.status_code == 400 and r.json()["error"]["code"] == "invalid_request"
        r = c.post("/api/analyze", content=b"{nope", headers={"Content-Type": "application/json"})
        assert r.status_code == 400
        r = c.post("/api/analyze", json={"text": 5})
        assert r.status_code == 400
        r = c.post("/api/analyze", json=["text"])
        assert r.status_code == 400
        r = c.get("/api/jobs/j_doesnotexist")
        assert r.status_code == 404 and r.json()["error"]["code"] == "job_not_found"
        r = c.get("/api/nothing-here")
        assert r.status_code == 404 and r.json()["error"]["code"] == "not_found"


def test_analyze_end_to_end(settings):
    (settings.calibration_dir / "thresholds.json").write_text(json.dumps(thresholds()))
    with make_client(settings, backend=FakeBackend(1.5)) as c:
        r = analyze(c, words(200))
        assert r.status_code == 202
        body = r.json()
        assert body["job_id"].startswith("j_") and body["position"] == 0
        snap = wait_job(c, body["job_id"])
        assert snap["status"] == "done" and snap["error"] is None and snap["position"] == 0
        res = snap["result"]
        assert res["verdict"]["band"] == "likely_ai"
        assert res["verdict"]["ai_percent"] == 100
        assert res["model_pair"] == PAIR
        assert set(res["timing_ms"]) == {"queue", "load", "score"}
        assert r.headers["cache-control"] == "no-store"
        # the submitted text is dropped once scored
        assert c.svc.queue._jobs[body["job_id"]].text == ""


def test_analyze_rate_limit(settings):
    lim = Limits(checks_per_hour=3)
    with make_client(settings, limits=lim) as c:
        a = {"CF-Connecting-IP": "203.0.113.9"}
        assert analyze(c, words(5), headers=a).status_code == 422  # rejected input does not count
        for _ in range(3):
            assert analyze(c, words(25), headers=a).status_code == 202
        r = analyze(c, words(25), headers=a)
        assert r.status_code == 429
        err = r.json()["error"]
        assert err["code"] == "rate_limited" and 3500 <= err["retry_after_s"] <= 3600
        assert r.headers["retry-after"] == str(err["retry_after_s"])
        assert analyze(c, words(25), headers={"CF-Connecting-IP": "203.0.113.10"}).status_code == 202
        # without CF-Connecting-IP the socket address is the key
        assert analyze(c, words(25)).status_code == 202


def test_queue_order_and_positions(settings):
    gate = threading.Event()
    log = []
    backend = FakeBackend(gate=gate, log=log)
    with make_client(settings, backend=backend) as c:
        texts = [words(20 + i, word=f"t{i}x") for i in range(3)]
        ids = []
        for i, t in enumerate(texts):
            r = analyze(c, t)
            assert r.status_code == 202
            ids.append(r.json()["job_id"])
            if i == 0:
                wait_job(c, ids[0], statuses=("scoring",))
        assert [c.get(f"/api/jobs/{j}").json()["position"] for j in ids] == [0, 1, 2]
        assert [c.get(f"/api/jobs/{j}").json()["status"] for j in ids] == ["scoring", "queued", "queued"]
        st = c.get("/api/status").json()
        assert st["queue_length"] == 3 and st["model_state"] == "ready"
        gate.set()
        snaps = [wait_job(c, j) for j in ids]
        assert all(s["status"] == "done" for s in snaps)
        assert [len(x) for x in log] == [20, 21, 22]  # scored in submission order
        assert c.get("/api/status").json()["queue_length"] == 0


def test_loading_model_state(settings):
    load_gate = threading.Event()
    with make_client(settings, backend=FakeBackend(load_gate=load_gate)) as c:
        assert c.get("/api/status").json()["model_state"] == "unloaded"
        jid = analyze(c, words(25)).json()["job_id"]
        assert wait_job(c, jid, statuses=("loading_model",))["status"] == "loading_model"
        assert c.get("/api/status").json()["model_state"] == "loading"
        load_gate.set()
        snap = wait_job(c, jid)
        assert snap["status"] == "done"
        assert c.get("/api/status").json()["model_state"] == "ready"


def test_queue_full(settings):
    gate = threading.Event()
    with make_client(settings, backend=FakeBackend(gate=gate), max_queue=1) as c:
        first = analyze(c, words(25)).json()["job_id"]
        wait_job(c, first, statuses=("scoring",))
        assert analyze(c, words(25)).status_code == 202
        r = analyze(c, words(25))
        assert r.status_code == 503 and r.json()["error"]["code"] == "busy"
        gate.set()


def test_idle_unload_and_reload(settings):
    with make_client(settings, idle_unload_s=0.2) as c:
        jid = analyze(c, words(25)).json()["job_id"]
        assert wait_job(c, jid)["status"] == "done"
        b1 = c.backends[0]
        assert b1.loaded
        assert wait_for(lambda: c.get("/api/status").json()["model_state"] == "unloaded", timeout=5)
        assert b1.unload_calls == 1 and not b1.loaded
        assert c.svc.manager.unloads == 1
        jid = analyze(c, words(25)).json()["job_id"]
        snap = wait_job(c, jid)
        assert snap["status"] == "done"
        assert c.svc.manager.loads == 2
    assert c.svc.manager.state == "unloaded"  # shutdown unloads too


def test_results_expire(settings):
    with make_client(settings, result_ttl_s=0.2) as c:
        jid = analyze(c, words(25)).json()["job_id"]
        assert wait_job(c, jid)["status"] == "done"
        assert wait_for(lambda: c.get(f"/api/jobs/{jid}").status_code == 404, timeout=5)


def test_job_errors(settings):
    class Broken(FakeBackend):
        def score_ids(self, ids):
            raise RuntimeError("boom")

    def failing_factory():
        raise OSError("no model file")

    with make_client(settings, factory=failing_factory) as c:
        snap = wait_job(c, analyze(c, words(25)).json()["job_id"])
        assert snap["status"] == "error" and snap["error"]["code"] == "model_unavailable"
        assert c.get("/api/status").json()["model_state"] == "unloaded"
    with make_client(settings, backend=Broken()) as c:
        snap = wait_job(c, analyze(c, words(25)).json()["job_id"])
        assert snap["status"] == "error" and snap["error"]["code"] == "scoring_failed"
        # the worker survives a failed job
        c.backends[0].score_ids = FakeBackend.score_ids.__get__(c.backends[0])
        assert wait_job(c, analyze(c, words(25)).json()["job_id"])["status"] == "done"
    with make_client(settings, limits=Limits(max_tokens=10)) as c:
        snap = wait_job(c, analyze(c, words(25)).json()["job_id"])
        assert snap["error"]["code"] == "too_long"


# ------------------------------------------------------------------------------------- status / accuracy


def test_status_and_accuracy(settings):
    with make_client(settings) as c:
        st = c.get("/api/status").json()
        assert st == {
            "model_state": "unloaded", "queue_length": 0, "model_pair": PAIR,
            "calibration": {"version": None, "ready": False},
            "limits": {"min_words": 20, "percent_min_words": 150, "max_words": 2500, "checks_per_hour": 30},
        }
        r = c.get("/api/accuracy")
        assert r.status_code == 404 and r.json()["error"]["code"] == "not_calibrated"

        report = {"version": "2026-10-07", "model_pair": PAIR, "held_out": {"human": 1, "ai": 2}}
        (settings.calibration_dir / "report.json").write_text(json.dumps(report))
        (settings.calibration_dir / "thresholds.json").write_text(json.dumps(thresholds()))
        assert c.get("/api/accuracy").json() == report
        assert c.get("/api/status").json()["calibration"] == {"version": "2026-10-07", "ready": True}

        other = thresholds(pair={"observer": "A", "performer": "B"})
        (settings.calibration_dir / "thresholds.json").write_text(json.dumps(other) + " ")
        assert c.get("/api/status").json()["calibration"] == {"version": "2026-10-07", "ready": False}


# ------------------------------------------------------------------------------------- extract


def _docx_bytes():
    import docx

    d = docx.Document()
    d.add_paragraph("First paragraph of the essay.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "cell one"
    t.cell(0, 1).text = "cell two"
    d.add_paragraph("Last paragraph — with “quotes”.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf_bytes(text: str | None):
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


def upload(c, name, data):
    return c.post("/api/extract", files={"file": (name, data, "application/octet-stream")})


def test_extract_formats(settings):
    with make_client(settings) as c:
        r = upload(c, "essay.txt", "Café text with “quotes”.\r\n\r\n\r\n\r\nSecond   para.\n".encode())
        assert r.status_code == 200
        assert r.json() == {"text": "Café text with “quotes”.\n\nSecond   para.", "filename": "essay.txt", "word_count": 6}
        assert upload(c, "old.TXT", "caf\xe9 au lait".encode("cp1252")).json()["text"] == "café au lait"
        assert upload(c, "u16.txt", "hello wide world".encode("utf-16")).json()["text"] == "hello wide world"

        r = upload(c, "paper.docx", _docx_bytes())
        assert r.status_code == 200
        assert r.json()["text"] == ("First paragraph of the essay.\n\ncell one\tcell two\n\n"
                                    "Last paragraph — with “quotes”.")

        r = upload(c, "scan.pdf", _pdf_bytes("Hello from a PDF file"))
        assert r.status_code == 200, r.text
        assert "Hello from a PDF file" in r.json()["text"]
        assert r.json()["filename"] == "scan.pdf"


def test_extract_errors(settings):
    with make_client(settings, max_upload_bytes=2000) as c:
        for name, data in (("pic.png", b"\x89PNG...."), ("noext", b"hello"), ("bad.docx", b"not a zip"),
                           ("bad.pdf", b"hello world")):
            r = upload(c, name, data)
            assert r.status_code == 415 and r.json()["error"]["code"] == "unsupported_file", name
        r = upload(c, "big.txt", b"a " * 1100)  # over max_upload_bytes, under the body cap
        assert r.status_code == 413 and r.json()["error"]["code"] == "file_too_large"
        r = upload(c, "huge.txt", b"a " * 50_000)  # over the request body cap
        assert r.status_code == 413 and r.json()["error"]["code"] == "file_too_large"
        for name, data in (("empty.txt", b"   \n\n "), ("blank.pdf", _pdf_bytes(None))):
            r = upload(c, name, data)
            assert r.status_code == 422 and r.json()["error"]["code"] == "no_text_found", name
        r = c.post("/api/extract", files={"other": ("a.txt", b"hello")})
        assert r.status_code == 400
        r = c.post("/api/extract", json={"text": "x"})
        assert r.status_code == 415


# ------------------------------------------------------------------------------------- web UI


def test_static_files_and_spa_fallback(settings, tmp_path):
    with make_client(settings) as c:
        r = c.get("/")
        assert r.status_code == 404 and "not built" in r.text

    dist = settings.web_dist
    (dist / "assets").mkdir(parents=True)
    (dist / "index.html").write_text("<!doctype html><title>Squint</title>")
    (dist / "assets" / "app-123.js").write_text("console.log(1)")
    (dist / "favicon.svg").write_text("<svg/>")
    (tmp_path / "secret.txt").write_text("top secret")
    with make_client(settings) as c:
        for path in ("/", "/check/abc", "/accuracy"):
            r = c.get(path)
            assert r.status_code == 200 and "Squint" in r.text and r.headers["cache-control"] == "no-cache"
        r = c.get("/assets/app-123.js")
        assert r.status_code == 200 and "immutable" in r.headers["cache-control"]
        assert "javascript" in r.headers["content-type"]
        assert c.get("/favicon.svg").headers["content-type"].startswith("image/svg")
        assert c.get("/assets/missing.js").status_code == 404
        for path in ("/..%2Fsecret.txt", "/assets/..%2F..%2Fsecret.txt", "/%2e%2e/secret.txt"):
            r = c.get(path)
            assert "top secret" not in r.text
        r = c.get("/")
        assert r.headers["x-content-type-options"] == "nosniff"
        assert r.headers["x-frame-options"] == "DENY"
